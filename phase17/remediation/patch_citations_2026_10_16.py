"""Targeted, logged citation corrections to `Methodology Draft.docx` (2026-10-16).

Resolves every reference the Methodology marked "CITATION NEEDS VERIFICATION". Each replacement was checked against
its primary source on 2026-10-16: the arXiv abstract records (1910.01526, 2312.14197, 2410.10813, 2504.19413,
2502.12110, 2111.09543), the AAAI proceedings page for Gated Linear Networks, the ACL Anthology records for SNLI
(D15-1075), MultiNLI (N18-1101) and Sentence-BERT (D19-1410), the Hugging Face model card of
cross-encoder/nli-deberta-v3-base, and the qdrant/qdrant GitHub repository. Venue fields are taken only from those
records (Mem0 lists no venue, so none is given).

Same rules as `patch_methodology_docx.py`: each paragraph must match exactly once or already be patched
(idempotent); only the citation text changes; every change is logged to
`remediation/logs/citations_patch_log.json`. The pre-patch file is preserved by git.
"""
from __future__ import annotations

import json
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
DOC = ROOT / "Methodology Draft.docx"
LOG = ROOT / "remediation" / "logs" / "citations_patch_log.json"

GLN = ('Veness, J., Lattimore, T., Budden, D., Bhoopchand, A., Mattern, C., Grabska-Barwinska, A., Sezener, E., Wang, J., '
       'Toth, P., Schmitt, S., and Hutter, M. "Gated Linear Networks." Proceedings of the AAAI Conference on Artificial '
       'Intelligence 35(11), 10015-10023, 2021. doi:10.1609/aaai.v35i11.17202 (arXiv:1910.01526). Implemented from scratch '
       'as the online-learning (GLN) component of the Phase 11 GNN-GLN hybrid (phase11/gln/).')
BIPIA = ('Yi, J., Xie, Y., Zhu, B., Kiciman, E., Sun, G., Xie, X., and Wu, F. "Benchmarking and Defending Against Indirect '
         'Prompt Injection Attacks on Large Language Models" (BIPIA). KDD 2025. doi:10.1145/3690624.3709179 '
         '(arXiv:2312.14197). One of the external injection corpora from which Phase 17\'s injection-shaped attack texts '
         'were drawn (with MINJA, AgentPoison, InjecAgent and AgentDojo).')
HEADER_OLD = ('Entries marked CITATION NEEDS VERIFICATION could not be confirmed against a primary source during this pass '
              'and must be resolved before submission rather than assumed.')
HEADER_NEW = ('[2026-10-16: every entry previously marked "CITATION NEEDS VERIFICATION" has since been verified against its '
              'primary source (arXiv record, AAAI proceedings, ACL Anthology, model card or repository); Qdrant is cited '
              'as software because its repository provides no academic citation.]')
LME = ('LongMemEval — Wu, D., Wang, H., Yu, W., Zhang, Y., Chang, K.-W., and Yu, D. "LongMemEval: Benchmarking Chat '
       'Assistants on Long-Term Interactive Memory." ICLR 2025 (arXiv:2410.10813). Dataset host used by this project: '
       'https://huggingface.co/datasets/xiaowu0162/longmemeval-cleaned (MIT license).')
MEM0 = ('Mem0 (mem0ai) — Chhikara, P., Khant, D., Aryan, S., Singh, T., and Yadav, D. "Mem0: Building Production-Ready '
        'AI Agents with Scalable Long-Term Memory." arXiv:2504.19413, 2025 (preprint; the arXiv record lists no venue). '
        'Used as the installed library mem0ai 2.0.19 (PyPI).')
AMEM = ('A-MEM (A-mem-sys) — Xu, W., Liang, Z., Mei, K., Gao, H., Tan, J., and Zhang, Y. "A-MEM: Agentic Memory for '
        'LLM Agents." NeurIPS 2025 (arXiv:2502.12110). Used as the cloned implementation '
        'https://github.com/WujiangXu/A-mem-sys, commit f303dfc71e07bdc787f4bc135d4cea328ae30e99.')
NLI = ('cross-encoder/nli-deberta-v3-base — a Sentence-Transformers cross-encoder (Reimers, N., and Gurevych, I. '
       '"Sentence-BERT: Sentence Embeddings using Siamese BERT-Networks." EMNLP-IJCNLP 2019, 3982-3992, '
       'doi:10.18653/v1/D19-1410) fine-tuned from microsoft/deberta-v3-base (He, P., Gao, J., and Chen, W. "DeBERTaV3: '
       'Improving DeBERTa using ELECTRA-Style Pre-Training with Gradient-Disentangled Embedding Sharing." ICLR 2023, '
       'arXiv:2111.09543) on SNLI (Bowman, S. R., Angeli, G., Potts, C., and Manning, C. D. "A large annotated corpus for '
       'learning natural language inference." EMNLP 2015, 632-642, doi:10.18653/v1/D15-1075) and MultiNLI (Williams, A., '
       'Nangia, N., and Bowman, S. R. "A Broad-Coverage Challenge Corpus for Sentence Understanding through Inference." '
       'NAACL-HLT 2018, 1112-1122, doi:10.18653/v1/N18-1101). The model card names no individual authors and recommends '
       'no citation, so the pinned checkpoint is cited by its Hugging Face identifier.')
QDRANT = ('Qdrant — Qdrant vector database, open-source software (Apache-2.0), https://github.com/qdrant/qdrant; client '
          'library qdrant-client 1.19.0 as pinned in this project. The repository provides no recommended academic '
          'citation, so it is cited as software.')

# (id, paragraph prefix that identifies the target, full replacement text OR (old substring, new substring))
CORRECTIONS = [
    ("CIT-GLN", 'Veness, J.', GLN),
    ("CIT-BIPIA", 'Yi, J.', BIPIA),
    ("CIT-HEADER", 'The following citations are verified directly', (HEADER_OLD, HEADER_NEW)),
    ("CIT-LME", 'LongMemEval —', LME),
    ("CIT-MEM0", 'Mem0 (mem0ai) —', MEM0),
    ("CIT-AMEM", 'A-MEM (A-mem-sys) —', AMEM),
    ("CIT-NLI", 'cross-encoder/nli-deberta-v3-base —', NLI),
    ("CIT-QDRANT", 'Qdrant —', QDRANT),
]


def _set_text(p, new: str) -> None:
    runs = p.runs
    if not runs:
        p.add_run(new)
        return
    runs[0].text = new
    for r in runs[1:]:
        r.text = ""


def apply() -> dict:
    import docx

    d = docx.Document(str(DOC))
    log = {"utc": time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime()), "doc": DOC.name, "changes": []}
    for cid, prefix, rep in CORRECTIONS:
        hits = [p for p in d.paragraphs if p.text.startswith(prefix)]
        if len(hits) != 1:
            raise SystemExit(f"{cid}: expected exactly one paragraph starting {prefix!r}, found {len(hits)}")
        p = hits[0]
        old = p.text
        if isinstance(rep, tuple):
            if rep[1] in old:
                log["changes"].append({"id": cid, "status": "already applied"})
                continue
            if old.count(rep[0]) != 1:
                raise SystemExit(f"{cid}: header sentence not found exactly once")
            new = old.replace(rep[0], rep[1])
        else:
            if old == rep:
                log["changes"].append({"id": cid, "status": "already applied"})
                continue
            new = rep
        _set_text(p, new)
        log["changes"].append({"id": cid, "status": "applied", "before": old, "after": new})
    remaining = sum("NEEDS VERIFICATION" in p.text for p in d.paragraphs if "previously marked" not in p.text)
    if remaining:
        raise SystemExit(f"{remaining} paragraph(s) still say NEEDS VERIFICATION; nothing saved")
    d.save(str(DOC))
    LOG.parent.mkdir(parents=True, exist_ok=True)
    runs = json.loads(LOG.read_text(encoding="utf-8")) if LOG.exists() else []       # append-only: every run is kept
    runs.append(log)
    LOG.write_text(json.dumps(runs, indent=1, ensure_ascii=False), encoding="utf-8")
    return {c["id"]: c["status"] for c in log["changes"]}


if __name__ == "__main__":
    print(apply())
