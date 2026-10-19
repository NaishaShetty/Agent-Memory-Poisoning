"""Targeted phase: what the EXISTING human annotation can and cannot support. No new annotation.

Owner's clarification (2026-10-13, in chat): annotator 1 = the owner; annotator 2 = the owner's friend, who labelled
on the owner's laptop -- which is why both workbooks were last saved by the owner's Office account.

This module records only what the files themselves show (workbook metadata, instructions, completeness, comments),
re-computes agreement with item-bootstrap intervals, and lists which project conclusions depend on the labels.
Output: data/remediation/annotation_provenance.json.
"""
from __future__ import annotations

import hashlib
import json
from pathlib import Path
from typing import Dict

import numpy as np

ROOT = Path(__file__).resolve().parents[2]
HL = ROOT / "docs" / "phase17" / "human_labels"
DATA = ROOT / "phase17" / "data"
OUT = DATA / "remediation" / "annotation_provenance.json"


def _kappa(a, b):
    a, b = np.asarray(a, bool), np.asarray(b, bool)
    po = float((a == b).mean())
    pa, pb = a.mean(), b.mean()
    pe = pa * pb + (1 - pa) * (1 - pb)
    return None if pe == 1 else (po - pe) / (1 - pe)


def _boot_kappa(a, b, n=10_000, seed=2026):
    rng = np.random.default_rng(seed)
    a, b = np.asarray(a, bool), np.asarray(b, bool)
    ks = []
    for _ in range(n):
        ix = rng.integers(0, len(a), len(a))
        k = _kappa(a[ix], b[ix])
        if k is not None:
            ks.append(k)
    return [round(float(np.quantile(ks, .025)), 3), round(float(np.quantile(ks, .975)), 3)] if ks else None


def run() -> Dict:
    import openpyxl

    from phase17.human_labels import KEY, XLSX, XLSX2, _labels

    meta = {}
    for tag, path in (("annotator1", XLSX), ("annotator2", XLSX2)):
        wb = openpyxl.load_workbook(path)
        p = wb.properties
        ins = "\n".join(" ".join(str(x) for x in r if x) for r in wb["Instructions"].iter_rows(values_only=True) if any(r))
        rows = list(wb["Label these"].iter_rows(values_only=True))[1:]
        meta[tag] = {"file": str(path.relative_to(ROOT)), "creator": p.creator, "last_modified_by": p.lastModifiedBy,
                     "created_utc": str(p.created), "modified_utc": str(p.modified),
                     "labelled_rows": sum(str(r[5] or "").strip().upper() in ("YES", "NO") for r in rows),
                     "comments": sum(bool(r[6]) and str(r[6]).strip() not in ("", "None") for r in rows),
                     "instructions_sha256": hashlib.sha256(ins.encode("utf-8")).hexdigest()[:16],
                     "instructions_mention_key_ban": "KEY_do_not_open" in ins,
                     "instructions_mention_first_sheet_ban": "labelling_sheet.xlsx" in ins,
                     "sheet_columns": list(list(wb["Label these"].iter_rows(values_only=True))[0])}
    key = {k["item"]: k for k in json.loads(KEY.read_text(encoding="utf-8"))}
    h1, h2 = _labels(XLSX), _labels(XLSX2)
    agree = {}
    for t in ("compliance", "harm", "mechanism", "ALL"):
        ix = [i for i in sorted(set(h1) & set(h2)) if t == "ALL" or key[i]["type"] == t]
        a, b = [h1[i] == "YES" for i in ix], [h2[i] == "YES" for i in ix]
        k = _kappa(a, b)
        agree[t] = {"n": len(ix), "agreement": round(sum(x == y for x, y in zip(a, b)) / len(ix), 3),
                    "kappa": None if k is None else round(k, 3), "kappa_item_bootstrap_ci95": _boot_kappa(a, b) if k is not None else None,
                    "yes_a1": sum(a), "yes_a2": sum(b)}
    judge = {}
    for t in ("compliance", "harm"):
        ix = [i for i in sorted(set(h1) & set(h2)) if key[i]["type"] == t]
        j = [bool(key[i]["judge"]) for i in ix]
        for who, lab in (("annotator1", h1), ("annotator2", h2)):
            h = [lab[i] == "YES" for i in ix]
            judge[f"J0_{t}_vs_{who}"] = {"kappa": round(_kappa(j, h), 3), "ci95": _boot_kappa(j, h), "n": len(ix)}
    out = {
        "owner_clarification": {"annotator1": "the project owner/author",
                                "annotator2": "the owner's friend; labelled on the owner's laptop",
                                "source": "owner statement in chat, 2026-10-13 (not recorded in the files themselves)"},
        "workbooks": meta,
        "timeline_from_files": {
            "annotator2_sheet_generated": "commit 4fc400dd, 2026-10-09 10:52 IST (blind sheet: no judge verdicts, no first-annotator labels)",
            "annotator2_sheet_last_saved": meta["annotator2"]["modified_utc"] + " UTC",
            "annotator2_labels_committed": "commit 601a4c06, 2026-10-09 12:30 IST"},
        "verifiable": [
            "both sheets contain 100/100 labelled items (same items, same order)",
            "annotator 2's sheet as generated contained no judge verdicts and no annotator-1 labels (sheet columns; commit 4fc400dd)",
            "written instructions told annotator 2 not to open the KEY file or annotator 1's sheet and not to discuss items until done",
            "annotator 2's labelling window: sheet generated 10:52 IST, last saved 12:23 IST on 2026-10-09",
            "annotator 1 (the owner) labelled 2026-10-07, before the second sheet existed"],
        "not_verifiable_from_records": [
            "whether annotator 2 complied with the instructions (the KEY file and annotator 1's sheet were on the same laptop)",
            "whether the owner was present or discussed items during labelling",
            "annotator 2's qualifications or prior exposure to the project",
            "independence of annotator 2 from the author (a friend is not independent of the author)"],
        "instructions_framing_note": "the instructions say each row is 'one judgement an AI judge already made; you are "
                                     "checking it' -- no verdict is shown, but the framing is not neutral",
        "agreement": agree, "judge_C0_against_each_annotator": judge,
        "what_depends_on_the_labels": {
            "compliance vote (C0+C1+C3)": "SELECTED on 35 items where both annotators agree (compliance_judge_fix, round 2)",
            "harm judge J2": "SELECTED on annotator 1's labels; annotator 2's harm labels are identical (ceiling 36/40 YES)",
            "net memory effect": "reported under C0 (not human-selected) AND the vote -- both positive, so not label-dependent",
            "defence detection / false-positive rates": "no judge and no human label involved",
            "defence ranking on working subsets": "identical under C0, vote, robust definitions (validity_final)"},
    }
    OUT.write_text(json.dumps(out, indent=2, default=str), encoding="utf-8")
    return out


if __name__ == "__main__":
    r = run()
    print(json.dumps({k: r[k] for k in ("workbooks", "agreement", "judge_C0_against_each_annotator")}, indent=1, default=str))
