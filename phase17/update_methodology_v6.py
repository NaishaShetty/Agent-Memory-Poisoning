"""Bring `Methodology Draft.docx` up to date with Defence v6 (2026-10-14/15). Idempotent: the block sits between its own
start marker and the heading of Section 33 (after the 2026-10-14 block) and is replaced wholesale on every run; the
conclusion status paragraph is replaced in place. Content comes from phase17/results_sections_v6.py (the same generator
as the Results Analysis), so the two documents cannot disagree. Log: remediation/logs/methodology_update_v6.json.
"""
from __future__ import annotations

import json
import re
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
DOC = ROOT / "Methodology Draft.docx"
LOG = ROOT / "remediation" / "logs" / "methodology_update_v6.json"
MARK = "[Section 32.71: Defence Programme v6 (2026-10-14/15), generated from the committed v6 artifacts.]"
END_MARK = "[End of the Defence v6 update.]"
STATUS = "[Status 2026-10-15, Defence v6]"


def _content():
    from phase17.formulas import rows as formula_rows
    from phase17.results_sections_v6 import blocks

    out = [("p", MARK)]
    out.append(("p", "Methods summary. v6 kept the frozen closure rubric and pre-registered every hypothesis, dataset, selection "
                     "rule and acceptance criterion before building new data (Phase 1); development choices used only "
                     "development owners or development roots; each held-out set was frozen by digest before scoring; "
                     "defects found after a run are kept as *_DEFECTIVE_* outputs and disclosed."))
    for kind, x in blocks():
        if kind == "h":
            out.append(("h2", "32.71 " + x.split(". ", 1)[1]))
        elif kind == "h2":
            out.append(("h3", "32.71." + x.split(" ", 1)[0].split(".", 1)[1] + " " + x.split(" ", 1)[1]))
        elif kind == "cap":
            out.append(("cap", x.replace("Table 23.", "Table 32.71.")))
        else:
            out.append((kind, x))
    fr = formula_rows()
    start = next(i for i, r in enumerate(fr) if r[0] == "Provenance ledger (PL) and source union")
    out.append(("table", [fr[0]] + fr[start:]))
    out.append(("cap", "Table 32.71.F. Formulas added by Defence v6 (the full set is in the Results Analysis, 'Definitions and "
                       "Formulas')."))
    out.append(("p", END_MARK))
    return out


def _status_text():
    d = ROOT / "phase17" / "data" / "defence_v6" / "closure_v6.json"
    if d.exists():
        c = json.loads(d.read_text(encoding="utf-8"))
        parts = "; ".join(f"{g.split(': ')[-1]} {v['before'] * 100:.0f}% -> {v['after'] * 100:.0f}%" for g, v in c["gaps"].items())
        return (f"{STATUS} After Defence v6 (Section 32.71) the evidence-based operational closure under the frozen rubric is: "
                f"{parts}. {c.get('summary', '')}")
    return (f"{STATUS} Defence v6 (Section 32.71) is in progress: provenance, claim-level lineage, claim fingerprints, powered "
            "propagation and lineage recovery are complete; the sleeper experiments and closure scoring are pending, and the "
            "closure scores of 2026-10-14 stand until then.")


def run() -> dict:
    from docx import Document
    from docx.shared import Inches

    from phase17.figures_v6 import run as draw

    figs = {k: ROOT / v for k, v in draw().items() if v}
    doc = Document(str(DOC))
    log = {"removed_previous_block": 0, "inserted": 0, "figures": 0}
    pars = doc.paragraphs
    anchor = next(p for p in pars if p.style.name.startswith("Heading 1") and p.text.startswith("33. Overall Conclusion"))
    body = anchor._p.getparent()
    start = next((p for p in pars if p.text == MARK), None)
    if start is not None:
        el = start._p
        while el is not None and el is not anchor._p:
            nxt = el.getnext()
            body.remove(el)
            log["removed_previous_block"] += 1
            el = nxt
    nums = [int(m.group(1)) for p in doc.paragraphs for m in [re.match(r"Figure (\d+) - ", p.text)] if m]
    fig_no, prev = (max(nums) if nums else 0), None
    styles = {s.name for s in doc.styles}
    for kind, payload in _content():
        if kind == "cap" and prev == "fig":
            fig_no += 1
            payload = f"Figure {fig_no} - {payload}"
        if kind == "fig":
            path = figs.get(payload)
            prev = "fig" if path and path.exists() else "nofig"
            if prev == "fig":
                par = anchor.insert_paragraph_before()
                par.add_run().add_picture(str(path), width=Inches(6.3))
                log["figures"] += 1
            continue
        if kind == "cap" and prev == "nofig":
            prev = kind
            continue
        prev = kind
        if kind == "h2":
            anchor.insert_paragraph_before(payload, style="Heading 2")
        elif kind == "h3":
            anchor.insert_paragraph_before(payload, style="Heading 3" if "Heading 3" in styles else "Heading 2")
        elif kind in ("p", "cap"):
            anchor.insert_paragraph_before(payload)
        elif kind == "table":
            t = doc.add_table(rows=len(payload), cols=len(payload[0]))
            try:
                t.style = "Table Grid"
            except Exception:  # noqa: BLE001
                pass
            for i, row in enumerate(payload):
                for j, v in enumerate(row):
                    t.cell(i, j).text = str(v)
            anchor._p.addprevious(t._tbl)
        log["inserted"] += 1
    # conclusion status (replaced in place on every run)
    txt = _status_text()
    ex = next((p for p in doc.paragraphs if p.text.startswith(STATUS)), None)
    if ex is not None:
        for r in ex.runs[1:]:
            r._r.getparent().remove(r._r)
        if ex.runs:
            ex.runs[0].text = txt
        else:
            ex.add_run(txt)
    else:
        prev14 = next((p for p in doc.paragraphs if p.text.startswith("[Status 2026-10-14]")), None)
        if prev14 is not None:
            from docx.oxml import OxmlElement
            el = OxmlElement("w:p")
            prev14._p.addnext(el)
            from docx.text.paragraph import Paragraph
            Paragraph(el, prev14._parent).add_run(txt)
        else:
            anchor.insert_paragraph_before(txt)
    log["status"] = txt[:160]
    doc.save(str(DOC))
    LOG.parent.mkdir(parents=True, exist_ok=True)
    LOG.write_text(json.dumps(log, indent=1), encoding="utf-8")
    return log


if __name__ == "__main__":
    print(run())
