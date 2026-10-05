# -*- coding: utf-8 -*-
"""Replace the configuration recommendation with one conditioned on the false-positive budget.

Blocker: B17/B18 were recommended on descriptive numbers while being the worst discriminators
on the only corpus carrying matched benign text. The frontier sweep settles it, and the deciding
fact is structural rather than a matter of calibration:

    B17's false-positive rate on hold-out 3 cannot go below 33.5% and B18's below 35.0%, at ANY
    threshold, because both exclude whenever B13 excludes and B13 alone is at 33.5% there.

So the unions are pinned above a third of benign traffic on injection-style text, and no
recalibration reaches the ~4% region where B14 and deberta operate. A single recommendation
cannot be right for both regimes, so the recommendation is now conditioned on the budget.
"""
from __future__ import annotations

import json
from pathlib import Path

import docx

DATA = Path("phase17/data")
FR = json.loads((DATA / "holdout3_frontier.json").read_text(encoding="utf-8"))
MF = json.loads((DATA / "holdout3_matched_fpr.json").read_text(encoding="utf-8"))
H3 = json.loads((DATA / "eval_holdout3.json").read_text(encoding="utf-8"))
BB = json.loads((DATA / "baseline_breakdown_round6.json").read_text(encoding="utf-8"))
PW = json.loads((DATA / "eval_rejection_powered.json").read_text(encoding="utf-8"))

rank = {e["scorer"]: e for e in MF["ranking_at_matched_dev_fpr"]}
u = {}
for c in ("B17", "B18"):
    cur = FR["curves"][c]
    p = min(cur, key=lambda x: abs(x["dev_target"] - 0.082))
    lo = min(cur, key=lambda x: x["fpr"])
    u[c] = {"at8": p, "floor": lo}

TEXT = [
    "RECOMMENDATION (round 6, final). No single configuration is right for every deployment "
    "here, and an earlier version of this document recommending one was wrong twice over -- "
    "first naming B13 unqualified, then the union configurations on numbers that were "
    "descriptive rather than held out. The recommendation is therefore conditioned on two "
    "things a deployment knows and this benchmark does not: which threat it faces, and what "
    "false-positive rate it can absorb.",

    f"The deciding constraint is structural, not a matter of tuning. B17 and B18 exclude "
    f"whenever B13 excludes, and B13 alone sits at "
    f"{100*H3['configs']['B13']['matched_benign_fpr']['rate']:.1f}% false positives on the "
    f"in-the-wild corpus, so sweeping their B14 component moves B17 only within "
    f"{100*u['B17']['floor']['fpr']:.1f}-36.0% and B18 within "
    f"{100*u['B18']['floor']['fpr']:.1f}-36.7%. No threshold reaches the ~4% region where B14 "
    f"and the strongest baseline operate. The unions are pinned above a third of benign traffic "
    f"on injection-style text and cannot be calibrated out of it.",

    "For MEMORY POISONING on ordinary memory traffic -- the threat this benchmark exists for -- "
    f"B13 remains the choice: best on this project's own attacks "
    f"({100*BB['aggregate_own']['B13']['rate']:.1f}% against B14's "
    f"{100*BB['aggregate_own']['B14']['rate']:.1f}% and deberta's "
    f"{100*BB['aggregate_own']['deberta_pi_v2@dev']['rate']:.1f}%), 1/1088 ordinary-benign false "
    "positives, 2/40 false flags on legitimate standing instructions, and cheaper per admitted "
    "memory than the content-only stack it supersedes. Supply it with the owner's conversation: "
    "corroboration lifts every founding family and costs nothing at all on user-authored "
    "content.",

    f"Compose it with the rejection ledger where an attacker can retry against the same scope. "
    f"That is the one configuration change with a large, powered effect: adaptive evasion "
    f"{100*PW['per_arm']['B15']['genuine_adaptive_evasion_of_initially_excluded']['rate']:.1f}% "
    f"to {100*PW['per_arm']['B16@0.90']['genuine_adaptive_evasion_of_initially_excluded']['rate']:.1f}%, "
    f"p={PW['paired_B16_vs_B15']['p_exact']:.3g}. Its radius carries a 29.7% ledger-poisoning "
    f"blast radius, which is the trade a deployment must accept deliberately.",

    f"For INJECTION-SHAPED text arriving in memory, the choice depends on the budget. Under a "
    f"tight false-positive budget, B14 at a dev-calibrated threshold "
    f"({100*rank['B14']['recall']:.1f}% recall at {100*rank['B14']['realised_fpr']:.1f}%) or an "
    f"off-the-shelf deberta pass ({100*rank['deberta_pi_v2']['recall']:.1f}% at "
    f"{100*rank['deberta_pi_v2']['realised_fpr']:.1f}%) are the only candidates; the unions "
    f"cannot operate there at all. Where a third of benign traffic can be quarantined for "
    f"review, B18 is the strongest detector measured here "
    f"({100*u['B18']['at8']['recall']:.1f}% recall at "
    f"{100*u['B18']['at8']['fpr']:.1f}%, discrimination "
    f"{100*u['B18']['at8']['youden']:+.1f} against deberta's "
    f"{100*rank['deberta_pi_v2']['recall_minus_fpr']:+.1f}). That is a review-queue posture, not "
    f"an admission-control one, and it should be described as such rather than as a default.",

    "B17 and B18 remain hold-out-informed: they were selected after reading hold-out results, "
    "and hold-out 3 tested them once. Their numbers are reported, their selection history is "
    "stated at every point of use, and they are NOT presented as a validated generalisation "
    "claim.",
]

for PATH, anchor_text, heading in (
    ("Methodology Draft.docx", "33. Overall Conclusion", "32.42 Recommendation, Conditioned on the False-Positive Budget"),
):
    doc = docx.Document(PATH)
    if any(heading in p.text for p in doc.paragraphs):
        print(f"{PATH}: already applied")
        continue
    concl = next(p for p in doc.paragraphs if p.text.startswith(anchor_text))
    BODY = doc.paragraphs[4].style
    h = concl.insert_paragraph_before(heading)
    h.style = doc.styles["Heading 2"]
    for t in TEXT:
        q = concl.insert_paragraph_before(t)
        q.style = BODY
    doc.save(PATH)
    print(f"{PATH}: added {heading}")

# results-analysis builder
rb = Path("phase17/build_results_analysis_docx.py")
s = rb.read_text(encoding="utf-8")
if "RECOMMENDATION_TEXT" not in s:
    s = s.replace('MHA = L("multihop_multiagent.json")',
                  'MHA = L("multihop_multiagent.json")\n'
                  f'RECOMMENDATION_TEXT = {TEXT!r}')
    s = s.replace('H("11. Limitations", 1)',
                  'H("10B. Recommendation, Conditioned on the False-Positive Budget", 1)\n'
                  'for _t in RECOMMENDATION_TEXT:\n    P(_t)\n\n'
                  'H("11. Limitations", 1)')
    rb.write_text(s, encoding="utf-8")
    print("results builder: recommendation added")
print("DONE")
