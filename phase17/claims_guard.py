"""Superseded claims must not stand unqualified where a reader meets them.

WHY THIS EXISTS (round 7). Round 6 corrected several of this project's own claims -- multi-hop is
an evaluator artifact rather than a reasoning ceiling, B14's instability is not "inherited from
deberta", adaptive evasion was resolved by a powered re-run, exfiltration is a tested mechanism.
Each correction was written as a NEW section, and the sentences it corrected were left standing
elsewhere, unqualified: the results paper's abstract still said adaptive evasion "remains open",
the methodology's overall conclusion still declared multi-hop outside the benchmark, the summary
file still printed a working-set rate (45.1%) that its own table contradicted, and two runbooks
still asserted that temperature-0 runs are bit-reproducible after the noise floor showed they are
not. A reader who meets the old sentence first has no way to know it was withdrawn.

Reporting a correction "in place" means the superseded wording either goes, or travels with its
correction. This module lists each superseded claim as a pattern, together with the qualifier
that must appear in the same paragraph if the wording is retained (for example in a section that
records the history). `violations(paragraphs)` returns every paragraph where a pattern occurs
without its qualifier.

It is applied to every generated document and summary file by
`phase17/tests/test_round7_discipline.py`, and the guard is proven by mutation testing in
`error_registry.py`: a superseded sentence is inserted into a COPY of each document and the test
must fail.
"""
from __future__ import annotations

import os
import re
from pathlib import Path
from typing import Iterable, List, Tuple

ROOT = Path(__file__).resolve().parents[1]

# (id, pattern of the superseded claim, qualifier that must accompany it, what superseded it)
SUPERSEDED: List[Tuple[str, str, str, str]] = [
    ("adaptive_open",
     r"adaptive evasion (remains|is|was recorded as) (open|unresolved)|adaptive evasion remains open",
     r"resolved|superseded|powered re-run|9\.3e-10|9\.31e-10|32\.39|section 9\.3",
     "the powered re-run: 32/130 -> 1/130, p=9.3e-10 (Methodology 32.39)"),
    ("multihop_scoped_out",
     r"multi-hop reasoning is (declared )?outside what this benchmark measures",
     r"correct(ed|ion)|evaluator artifact|withdrawn|superseded|32\.34",
     "multi-hop is an NLI-strictness artifact (Methodology 32.34)"),
    ("deberta_inherited",
     r"inherited from deberta|instability (is )?inherited",
     r"imprecise|correct(ed|ion)|superseded|ablation|32\.47",
     "removing deberta widens the gap; the feature is load-bearing (Methodology 32.47)"),
    ("working_rate_45",
     r"45\.1%",
     r"CORRECTED",
     "the working-set rate is 42.5% (1,117/2,628)"),
    ("batching_bit_reproducible",
     r"batching changes throughput only and never a result|results are unaffected either way|"
     r"a resumed result is identical to an uninterrupted one",
     r"noise floor|not bit-reproducible|5 points",
     "identical prompts scored 66, 64, 61 of 120: Ollama batching is not bit-reproducible"),
    ("exfil_untested",
     r"exfiltration[_ a-z]* (is|must be reported as) an? untested mechanism",
     r"finally tested|now a tested|supersed|32\.43",
     "exfiltration is a tested mechanism; the negative is about the attack class (32.43)"),
    ("exfil_weakest_authoring",
     r"exfiltration remains the weakest family",
     r"tested|supersed|32\.43",
     "superseded by the structural-gate regeneration and the three-agent test (32.43)"),
    ("b13_recommendation_unqualified",
     r"B13 remains the recommendation",
     r"memory poisoning|conditioned|false-positive budget",
     "the recommendation is conditioned on threat and false-positive budget (32.42/32.45)"),
    ("dev_snapshot_disjoint",
     r"(same )?disjoint dev snapshot|asserted disjoint from the frozen",
     r"R7-1|not disjoint|copies of|contaminat|hold-out-calibrated",
     "382/400 'dev' items are hold-out 3 prompts under a different wrapper (round 7, R7-1)"),
    ("holdout3_dev_calibrated",
     r"dev-calibrated threshold \(38\.6%|matched dev-calibrated|calibrated on (the )?dev negatives",
     r"hold-out-calibrated|R7-1|32\.49|10C|near-copies",
     "the hold-out 3 'dev' calibration set was the hold-out's own benign half (R7-1)"),
    ("ledger_no_effect_conclusion",
     r"showed no significant effect at its frozen radius",
     r"powered|resolved|9\.3e-10|32\.39|section 9\.3",
     "superseded by the powered re-run (32.39)"),
]


def _split(text: str) -> List[str]:
    # Markdown wraps sentences across lines, so whitespace is collapsed within a paragraph, or a
    # claim split over a line break would be invisible to the patterns.
    paras = re.split(r"\n\s*\n|\n(?=#)|\n(?=\|)", text)
    return [" ".join(p.split()) for p in paras if p.strip()]


def violations(paragraphs: Iterable[str]) -> List[Tuple[str, str, str]]:
    """(claim id, offending paragraph excerpt, what superseded it) for every unqualified hit."""
    out = []
    for para in paragraphs:
        for cid, pat, qual, why in SUPERSEDED:
            if re.search(pat, para, re.I) and not re.search(qual, para, re.I):
                m = re.search(pat, para, re.I)
                s = max(0, m.start() - 80)
                out.append((cid, para[s:m.end() + 80].replace("\n", " "), why))
    return out


def docx_paragraphs(path: Path) -> List[str]:
    import docx

    d = docx.Document(str(path))
    paras = [p.text for p in d.paragraphs]
    for t in d.tables:
        for r in t.rows:
            paras.append(" | ".join(c.text for c in r.cells))
    return paras


def targets() -> List[Tuple[str, Path]]:
    """Every document a reader meets. An env override lets the mutation harness point a target
    at a tampered COPY, so the real document is never edited to prove the guard works."""
    base = [("results_docx", ROOT / "MAMBench Results Analysis.docx"),
            ("methodology_docx", ROOT / "Methodology Draft.docx"),
            ("final_results_md", ROOT / "docs" / "phase17" / "PHASE17_FINAL_RESULTS.md"),
            ("running_md", ROOT / "phase17" / "RUNNING_THE_PIPELINE.md"),
            ("resume_md", ROOT / "phase17" / "RESUME_AFTER_SHUTDOWN.md")]
    out = []
    for name, p in base:
        ov = os.environ.get(f"MAMBENCH_CLAIMS_{name.upper()}")
        out.append((name, Path(ov) if ov else p))
    return out


def paragraphs_of(path: Path) -> List[str]:
    if path.suffix == ".docx":
        return docx_paragraphs(path)
    return _split(path.read_text(encoding="utf-8"))
