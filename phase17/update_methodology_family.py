# -*- coding: utf-8 -*-
"""Add Section 32.32 (per-family results, named) and 32.33 (the corroboration floor) to the
methodology, in response to a direct question: why are FARMA, DSRM, MemoryGraft and MINJA not
visible anywhere in the reporting?

They were measured all along. Both documents organised every detection table by CORPUS, which
is how the populations were frozen, so the seven founding families were spread across them and
appeared by name only in the harm table. Worse, the evaluation those tables come from never
supplies corroboration evidence, so the provenance configuration's strongest signal is inert in
all of them and its numbers are floors.
"""
from __future__ import annotations

import json
from pathlib import Path

import docx

PATH = "Methodology Draft.docx"
DATA = Path("phase17/data")
MARKER = "32.32 Results by Attack Family"


def L(n):
    p = DATA / n
    return json.loads(p.read_text(encoding="utf-8")) if p.exists() else {}


FAM, EVID5, EVID = L("family_breakdown.json"), L("evidence_available_check.json"), L("evidence_round6.json")

doc = docx.Document(PATH)
if any(MARKER in p.text for p in doc.paragraphs):
    raise SystemExit("already applied")
concl = next(p for p in doc.paragraphs if p.text.startswith("33. Overall Conclusion"))
BODY = doc.paragraphs[4].style


def h2(t):
    p = concl.insert_paragraph_before(t)
    p.style = doc.styles["Heading 2"]
    return p


def body(t):
    p = concl.insert_paragraph_before(t)
    p.style = BODY
    return p


def table(rows):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    try:
        t.style = "Table Grid"
    except KeyError:
        pass
    for i, r in enumerate(rows):
        for j, c in enumerate(r):
            cell = t.cell(i, j)
            cell.text = "" if c is None else str(c)
            if i == 0:
                for run in cell.paragraphs[0].runs:
                    run.bold = True
    concl._p.addprevious(t._tbl)
    body("")
    return t


h2(MARKER)
body(
    "Every detection table in Sections 32.19-32.31 is organised by CORPUS, because that is the "
    "unit in which the populations were built, frozen and digested. The seven founding families "
    "this benchmark reconstructs from the literature -- AgentPoison, MINJA, FARMA, MemoryGraft, "
    "DSRM, MPBench-PCFI and sleeper memory poisoning -- are distributed across those corpora, "
    "with the consequence that a reader looking for a named family's detection rate could find "
    "it in the harm study and nowhere else. `phase17/family_breakdown.py` joins the existing "
    "artifacts and reports each family by name; no model is re-run, because the measurements "
    "existed and only the view was missing.")
if FAM:
    rows = [["Attack family", "n", "Mechanism-correct", "Working rate",
             "B12 full/working", "B13 full/working", "B14 full/working", "B17 full/working"]]

    def c(v, cfg):
        if cfg not in v:
            return "-"
        f = v[cfg]["full"]["rate"]
        w = v[cfg]["working"]["rate"] if v[cfg].get("working") else None
        return ("%.0f%% / " % (100 * f)) + (("%.0f%%" % (100 * w)) if w is not None else "-")

    for group in (FAM["founding_families"], FAM["mechanism_families"]):
        for f in group:
            v = FAM["families"].get(f)
            if not v:
                rows.append([f, "absent", "-", "-", "-", "-", "-", "-"])
                continue
            mc = v.get("mechanism_correct")
            rows.append([f + (" *" if v["is_founding"] else ""), v["n_instances"],
                         ("%.0f%%" % (100 * mc["rate"])) if mc else "-",
                         "%.0f%%" % (100 * v["working_rate"]),
                         c(v, "B12"), c(v, "B13"), c(v, "B14"), c(v, "B17")])
    table(rows)
    body(
        "Table 3 - Per-family results; * marks the seven founding families. Of those seven only "
        "FARMA (74% full, 82% on working attacks) and AgentPoison (59%, 86% working, under the "
        "fusion configurations) are well detected. DSRM, MPBench-PCFI, MINJA and MemoryGraft sit "
        "between 11% and 32%, and sleeper memory poisoning is at 0% under every configuration "
        "including on its 24 working instances. The later mechanism families are detected far "
        "better -- policy revocation 98%, authority impersonation 82%, conditional backdoor 52% "
        "-- which is the expected shape, since those are the classes the defense stack was "
        "designed against while the founding families are the independent test of it. The "
        "detection columns are nonetheless a FLOOR, for the reason given in Section 32.33.")
    body(
        "The sleeper figure has an identifiable cause rather than being a brute detection "
        "failure, and the cause matters for the threat model. That attack is specified as a "
        "mundane, harmless-looking statement about a person's routine containing no instruction, "
        "and its mechanism check requires precisely that; the family is 100% mechanism-correct "
        "and 63% working, so the instances are valid and they do function. A content-only "
        "detector that flagged such a note would also flag the benign personal notes that make "
        "up most of a real memory store, so a 0% content-only rate on this family is the attack "
        "succeeding at what it was constructed to do. It is also exactly the class that "
        "provenance and corroboration exist to reach.")

h2("32.33 Corroboration Was Inert in the Round-6 Evaluation, and Its Numbers Are Floors")
body(
    "The B12 and B13 columns of Table 3 are identical for all seven founding families, which is "
    "not a coincidence but a property of the harness. `eval_round6` supplies provenance SOURCES "
    "and never supplies corroboration EVIDENCE, so B13's NLI check of personal claims against "
    "the owner's own conversation turns -- the signal that reaches additive fabrication at 37/72 "
    "in the pre-registered comparison of Section 32.24 -- cannot fire. Every B13, B14, B15 and "
    "B17 figure in the corpus-oriented tables of this document is therefore a floor for the "
    "provenance-aware configurations. Round 5 established this and labelled its own numbers as "
    "floors; the label was not carried into per-family reporting, where its absence reads as a "
    "flat detection failure. This is a defect in the reporting, not in the defense.")
if EVID5:
    rows = [["Population", "Config", "No evidence", "With the real conversation", "Exact test"]]
    for pop, v in EVID5["populations"].items():
        for cfg, x in v.items():
            if isinstance(x, dict) and "no_evidence" in x:
                rows.append([pop, cfg, f"{x['no_evidence']['k']}/{x['no_evidence']['n']}",
                             f"{x['with_evidence']['k']}/{x['with_evidence']['n']}",
                             "p = %.4g" % x["paired"]["p_exact"]])
    table(rows)
    body(
        f"Table 4 - The magnitude of what is switched off, on the three populations whose real "
        f"conversation exists. The evidence is {EVID5.get('n_evidence_turns')} turns of the "
        f"actual owner conversation (speakers "
        f"{', '.join(EVID5.get('speakers', []))}), not material constructed for the test, and the "
        f"poison is marked as arriving on an untrusted channel exactly as the threat model "
        f"states. On the 47 defense-aware evasion rewrites -- stuck at 4/47 for every "
        f"configuration across three rounds and reported as the ceiling of content-only "
        f"detection -- supplying the conversation takes B13 to 28/47 (p<0.0001), the largest "
        f"single detection improvement in this project, obtained by making a signal available "
        f"rather than by changing a detector.")
body(
    "`phase17/evidence_round6.py` extends this to every own-attack population and family, with "
    "the condition that makes the comparison honest: THE SAME EVIDENCE IS SUPPLIED TO THE BENIGN "
    "POPULATIONS. Supplying an owner conversation renders every claim about a person who does not "
    "appear in it uncorroborated, which is genuine signal against an invented-persona attack and "
    "equally a genuine risk to a benign note about someone the conversation never mentions, so a "
    "recall gain reported without its false-positive cost measured on the same footing would be "
    "meaningless. Both arms are run and both are reported; if the false-positive rate rises as "
    "much as recall, the artifact records that there is no gain."
    + ("" if EVID else " That run had not completed when this revision was generated."))
body(
    "What this does not fix, and the boundary is sharp: the external injection corpora have no "
    "owner conversation, and their attacks are not claims about a person, so there is nothing to "
    "corroborate. Evidence is inapplicable to them by construction. The provenance "
    "configuration's weak showing on those corpora is therefore genuine rather than an artifact, "
    "and it is the reason the union configurations of Section 32.28 exist at all.")

doc.save(PATH)
print("added Sections 32.32 and 32.33 to the methodology")
print("DONE")
